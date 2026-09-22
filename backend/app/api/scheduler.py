"""초기 생산 스케줄러 API."""

from calendar import monthrange
from datetime import date
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.dependencies import get_db
from app.models.schedule import ProductionScheduleItem, RawMaterialDailyBalance, RawMaterialValidationRun, ScheduleChangeHistory, ScheduleRun, SchedulerSetting, UnscheduledRequirement, WorkCalendarDay
from app.models.planning_master import Plant, ProductionLine
from app.models.planning_run import PlanningRun
from app.schemas.schedule import CalendarDayResponse, CalendarSettingsResponse, CalendarSettingsUpdate, RawMaterialDailyBalanceResponse, RawMaterialValidationResponse, ScheduleCopyRequest, ScheduleCreateRequest, ScheduleItemCreate, ScheduleItemUpdate, ScheduleRunResponse, ScheduleStatusUpdate, ProductionScheduleItemResponse, UnscheduledRequirementResponse
from app.services.scheduler import create_schedule_run, update_schedule_item, reschedule_shortfalls
from app.services.raw_material_validator import validate_schedule_raw_materials

router = APIRouter(prefix="/scheduler", tags=["scheduler"])


def _response(run: ScheduleRun, items: list[ProductionScheduleItem], shortages: list[UnscheduledRequirement]) -> ScheduleRunResponse:
    return ScheduleRunResponse(
        id=run.id,
        planning_run_id=run.planning_run_id,
        created_at=run.created_at,
        items=[ProductionScheduleItemResponse.model_validate(item) for item in items],
        shortages=[UnscheduledRequirementResponse.model_validate(item) for item in shortages],
        version=run.version, status=run.status, change_reason=run.change_reason, confirmed_at=run.confirmed_at,
    )


@router.post("/runs/{planning_run_id}", response_model=ScheduleRunResponse, status_code=status.HTTP_201_CREATED)
def create_run(planning_run_id: int, payload: ScheduleCreateRequest, db: Session = Depends(get_db)) -> ScheduleRunResponse:
    try:
        run, items, shortages = create_schedule_run(db, planning_run_id, payload.conditions)
        return _response(run, items, shortages)
    except ValueError as error:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(error)) from error


@router.get("/runs/versions")
def list_versions(db: Session = Depends(get_db)):
    rows = db.execute(select(ScheduleRun, PlanningRun).join(PlanningRun, PlanningRun.id == ScheduleRun.planning_run_id).order_by(ScheduleRun.created_at.desc())).all()
    return [{"id": run.id, "version": run.version, "created_at": run.created_at, "status": "저장" if run.status == "검토 중" else run.status, "planning_year": plan.year, "planning_month": plan.month, "planning_end_year": plan.end_year, "planning_end_month": plan.end_month} for run, plan in rows]


@router.get("/runs")
def list_runs(status: str | None = None, db: Session = Depends(get_db)):
    """상태별로 불러오기 화면에 표시할 생산계획 이력을 반환한다."""
    query = select(ScheduleRun, PlanningRun).join(PlanningRun, PlanningRun.id == ScheduleRun.planning_run_id)
    if status:
        # 이전 화면에서 사용하던 '검토 중'은 저장 이력으로 계속 조회할 수 있게 한다.
        query = query.where(ScheduleRun.status.in_(("저장", "검토 중")) if status == "저장" else ScheduleRun.status == status)
    rows = db.execute(query.order_by(ScheduleRun.created_at.desc())).all()
    return [{"id": run.id, "version": run.version, "created_at": run.created_at, "status": "저장" if run.status == "검토 중" else run.status, "planning_year": plan.year, "planning_month": plan.month, "planning_end_year": plan.end_year, "planning_end_month": plan.end_month} for run, plan in rows]


@router.get("/runs/compare")
def compare_versions(base_id: int, compare_id: int, db: Session = Depends(get_db)):
    base_run, compare_run = db.get(ScheduleRun, base_id), db.get(ScheduleRun, compare_id)
    if not base_run or not compare_run:
        raise HTTPException(status_code=404, detail="비교할 생산계획 버전을 찾을 수 없습니다.")
    process_by_line = {line.code: line.process_code for line in db.scalars(select(ProductionLine))}
    def aggregate(run_id: int):
        values: dict[tuple[str, str, str, str], float] = {}
        for item in db.scalars(select(ProductionScheduleItem).where(ProductionScheduleItem.schedule_run_id == run_id)):
            key = (item.planned_date.strftime("%Y-%m"), item.plant_name, item.product_code, process_by_line.get(item.line_code, "미분류"))
            values[key] = values.get(key, 0.0) + item.planned_quantity_ton
        return values
    base_values, compare_values = aggregate(base_id), aggregate(compare_id)
    keys = sorted(set(base_values) | set(compare_values))
    return {"base_version": {"id": base_run.id, "version": base_run.version, "created_at": base_run.created_at}, "compare_version": {"id": compare_run.id, "version": compare_run.version, "created_at": compare_run.created_at}, "items": [{"period": period, "plant_name": plant, "product_code": product, "process_code": process, "base_quantity_ton": base_values.get((period, plant, product, process), 0.0), "compare_quantity_ton": compare_values.get((period, plant, product, process), 0.0), "difference_ton": compare_values.get((period, plant, product, process), 0.0) - base_values.get((period, plant, product, process), 0.0)} for period, plant, product, process in keys]}


@router.get("/runs/latest", response_model=ScheduleRunResponse)
def get_latest_run(db: Session = Depends(get_db)) -> ScheduleRunResponse:
    """독립 화면에서도 마지막으로 만든 생산계획을 조회한다."""
    run = db.scalar(select(ScheduleRun).order_by(ScheduleRun.created_at.desc()))
    if not run:
        raise HTTPException(status_code=404, detail="생성된 생산 스케줄이 없습니다.")
    items = list(db.scalars(select(ProductionScheduleItem).where(ProductionScheduleItem.schedule_run_id == run.id).order_by(ProductionScheduleItem.planned_date, ProductionScheduleItem.plant_name, ProductionScheduleItem.line_code)))
    shortages = list(db.scalars(select(UnscheduledRequirement).where(UnscheduledRequirement.schedule_run_id == run.id).order_by(UnscheduledRequirement.product_code)))
    return _response(run, items, shortages)


@router.get("/runs/{schedule_run_id}", response_model=ScheduleRunResponse)
def get_run(schedule_run_id: int, db: Session = Depends(get_db)) -> ScheduleRunResponse:
    run = db.get(ScheduleRun, schedule_run_id)
    if not run:
        raise HTTPException(status_code=404, detail="스케줄 실행 이력을 찾을 수 없습니다.")
    items = list(db.scalars(select(ProductionScheduleItem).where(ProductionScheduleItem.schedule_run_id == schedule_run_id).order_by(ProductionScheduleItem.planned_date, ProductionScheduleItem.line_code)))
    shortages = list(db.scalars(select(UnscheduledRequirement).where(UnscheduledRequirement.schedule_run_id == schedule_run_id).order_by(UnscheduledRequirement.product_code)))
    return _response(run, items, shortages)


@router.get("/runs/{schedule_run_id}/change-history")
def get_change_history(schedule_run_id: int, db: Session = Depends(get_db)):
    if not db.get(ScheduleRun, schedule_run_id):
        raise HTTPException(status_code=404, detail="생산계획 실행 이력을 찾을 수 없습니다.")
    rows = list(db.scalars(
        select(ScheduleChangeHistory)
        .where(ScheduleChangeHistory.schedule_run_id == schedule_run_id)
        .order_by(ScheduleChangeHistory.changed_at.desc())
    ))
    return [{
        "id": row.id, "changed_at": row.changed_at, "changed_by": row.changed_by,
        "change_reason": row.change_reason, "before_values": row.before_values,
        "after_values": row.after_values,
    } for row in rows]


@router.patch("/items/{item_id}", response_model=ScheduleRunResponse)
def update_item(item_id: int, payload: ScheduleItemUpdate, db: Session = Depends(get_db)) -> ScheduleRunResponse:
    try:
        run, items, shortages = update_schedule_item(
            db,
            item_id,
            payload.planned_date,
            payload.product_code,
            payload.planned_quantity_ton,
            payload.downtime_hours,
            payload.work_rate,
            payload.operation_status,
            payload.auto_calculate,
            payload.is_locked,
            payload.adjustment_note,
        )
        return _response(run, items, shortages)
    except ValueError as error:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(error)) from error


@router.post("/runs/{schedule_run_id}/items", response_model=ScheduleRunResponse, status_code=status.HTTP_201_CREATED)
def create_item(schedule_run_id: int, payload: ScheduleItemCreate, db: Session = Depends(get_db)) -> ScheduleRunResponse:
    run = db.get(ScheduleRun, schedule_run_id)
    line = db.get(ProductionLine, payload.line_code)
    if not run or not line:
        raise HTTPException(status_code=404, detail="생산계획 또는 생산라인을 찾을 수 없습니다.")
    if run.status == "확정":
        raise HTTPException(status_code=400, detail="확정된 스케줄은 수정할 수 없습니다. 새 버전을 복사해 수정해 주세요.")
    if db.scalar(select(ProductionScheduleItem).where(ProductionScheduleItem.schedule_run_id == run.id, ProductionScheduleItem.planned_date == payload.planned_date, ProductionScheduleItem.line_code == line.code)):
        raise HTTPException(status_code=400, detail="같은 라인·날짜에 이미 생산계획이 있습니다.")
    plant = db.get(Plant, line.plant_code)
    item = ProductionScheduleItem(schedule_run_id=run.id, planned_date=payload.planned_date, plant_name=plant.name if plant else line.plant_code, line_code=line.code, line_name=line.name, product_code=payload.product_code, planned_quantity_ton=0, available_capacity_ton=0, downtime_hours=payload.downtime_hours or 0, work_rate=(payload.work_rate or 100) / 100, yield_rate=1, operation_status=payload.operation_status or "가동", changeover_hours=0, is_locked=payload.is_locked, adjustment_note=payload.adjustment_note)
    db.add(item); db.flush()
    try:
        updated_run, items, shortages = update_schedule_item(db, item.id, payload.planned_date, payload.product_code, payload.planned_quantity_ton, payload.downtime_hours, payload.work_rate, payload.operation_status, True, payload.is_locked, payload.adjustment_note)
        return _response(updated_run, items, shortages)
    except ValueError as error:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(error)) from error

@router.put("/runs/{schedule_run_id}/status", response_model=ScheduleRunResponse)
def update_status(schedule_run_id: int, payload: ScheduleStatusUpdate, db: Session = Depends(get_db)) -> ScheduleRunResponse:
    from datetime import datetime
    run=db.get(ScheduleRun,schedule_run_id)
    if not run or payload.status not in {"작성 중", "저장", "검토 중", "확정"}: raise HTTPException(400,"유효한 스케줄 상태가 아닙니다.")
    run.status=payload.status; run.change_reason=payload.change_reason.strip(); run.confirmed_at=datetime.now() if payload.status=="확정" else None; db.commit()
    items=list(db.scalars(select(ProductionScheduleItem).where(ProductionScheduleItem.schedule_run_id==run.id).order_by(ProductionScheduleItem.planned_date,ProductionScheduleItem.line_code)))
    shortages=list(db.scalars(select(UnscheduledRequirement).where(UnscheduledRequirement.schedule_run_id==run.id).order_by(UnscheduledRequirement.product_code)))
    return _response(run,items,shortages)

@router.post("/runs/{schedule_run_id}/copy", response_model=ScheduleRunResponse, status_code=status.HTTP_201_CREATED)
def copy_run(schedule_run_id: int, payload: ScheduleCopyRequest, db: Session = Depends(get_db)) -> ScheduleRunResponse:
    source=db.get(ScheduleRun,schedule_run_id)
    if not source: raise HTTPException(404,"복사할 생산 스케줄을 찾을 수 없습니다.")
    version=max([row.version for row in db.scalars(select(ScheduleRun).where(ScheduleRun.planning_run_id==source.planning_run_id))],default=0)+1
    run=ScheduleRun(planning_run_id=source.planning_run_id,version=version,status="작성 중",change_reason=payload.change_reason.strip()); db.add(run); db.flush()
    source_items=list(db.scalars(select(ProductionScheduleItem).where(ProductionScheduleItem.schedule_run_id==source.id)))
    items=[]
    for row in source_items:
        item=ProductionScheduleItem(schedule_run_id=run.id,planned_date=row.planned_date,plant_name=row.plant_name,line_code=row.line_code,line_name=row.line_name,product_code=row.product_code,planned_quantity_ton=row.planned_quantity_ton,available_capacity_ton=row.available_capacity_ton,downtime_hours=row.downtime_hours,work_rate=row.work_rate,yield_rate=row.yield_rate,operation_status=row.operation_status,changeover_hours=row.changeover_hours,is_locked=False,adjustment_note="")
        db.add(item); items.append(item)
    shortages=[]
    for row in db.scalars(select(UnscheduledRequirement).where(UnscheduledRequirement.schedule_run_id==source.id)):
        short=UnscheduledRequirement(schedule_run_id=run.id,product_code=row.product_code,required_quantity_ton=row.required_quantity_ton,scheduled_quantity_ton=row.scheduled_quantity_ton,unallocated_quantity_ton=row.unallocated_quantity_ton); db.add(short); shortages.append(short)
    db.commit(); [db.refresh(x) for x in items+shortages]; db.refresh(run); return _response(run,items,shortages)


@router.post("/runs/{schedule_run_id}/reschedule", response_model=ScheduleRunResponse)
def reschedule(schedule_run_id: int, db: Session = Depends(get_db)) -> ScheduleRunResponse:
    try:
        run, items, shortages, _ = reschedule_shortfalls(db, schedule_run_id)
        return _response(run, items, shortages)
    except ValueError as error:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(error)) from error


def _validation_response(run: RawMaterialValidationRun, balances: list[RawMaterialDailyBalance]) -> RawMaterialValidationResponse:
    return RawMaterialValidationResponse(
        id=run.id,
        schedule_run_id=run.schedule_run_id,
        created_at=run.created_at,
        balances=[RawMaterialDailyBalanceResponse.model_validate(balance) for balance in balances],
    )


@router.post("/runs/{schedule_run_id}/raw-material-validation", response_model=RawMaterialValidationResponse, status_code=status.HTTP_201_CREATED)
def validate_raw_materials(schedule_run_id: int, db: Session = Depends(get_db)) -> RawMaterialValidationResponse:
    try:
        run, balances = validate_schedule_raw_materials(db, schedule_run_id)
        return _validation_response(run, balances)
    except ValueError as error:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(error)) from error


def _calendar_response(db: Session, year: int, month: int) -> CalendarSettingsResponse:
    dates = [date(year, month, day) for day in range(1, monthrange(year, month)[1] + 1)]
    saved = {item.calendar_date: item for item in db.scalars(select(WorkCalendarDay).where(WorkCalendarDay.calendar_date.in_(dates)))}
    days = [
        CalendarDayResponse(
            calendar_date=day,
            is_working=saved[day].is_working if day in saved else True,
            note=saved[day].note if day in saved else "기본 24시간 가동",
        )
        for day in dates
    ]
    setting = db.get(SchedulerSetting, 1)
    return CalendarSettingsResponse(changeover_hours=setting.changeover_hours if setting else 0.0, days=days)


@router.get("/calendar", response_model=CalendarSettingsResponse)
def get_calendar(year: int, month: int, db: Session = Depends(get_db)) -> CalendarSettingsResponse:
    if not 2000 <= year <= 2100 or not 1 <= month <= 12:
        raise HTTPException(status_code=400, detail="유효한 연도와 월을 입력해 주세요.")
    return _calendar_response(db, year, month)


@router.put("/calendar", response_model=CalendarSettingsResponse)
def update_calendar(payload: CalendarSettingsUpdate, db: Session = Depends(get_db)) -> CalendarSettingsResponse:
    target_days = [date(payload.year, payload.month, day) for day in range(1, monthrange(payload.year, payload.month)[1] + 1)]
    holidays = set(payload.non_working_dates)
    for current_day in target_days:
        is_working = True
        note = "기본 24시간 가동"
        if current_day in holidays:
            is_working, note = False, "사용자 지정 휴일"
        existing = db.get(WorkCalendarDay, current_day)
        if existing:
            existing.is_working, existing.note = is_working, note
        else:
            db.add(WorkCalendarDay(calendar_date=current_day, is_working=is_working, note=note))
    setting = db.get(SchedulerSetting, 1)
    if setting:
        setting.changeover_hours = payload.changeover_hours
    else:
        db.add(SchedulerSetting(id=1, changeover_hours=payload.changeover_hours))
    db.commit()
    return _calendar_response(db, payload.year, payload.month)


def _write_sheet(sheet, title: str, headers: list[str], rows: list[list[object]]) -> None:
    sheet.title = title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1769AA")
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(28, max(12, max(len(str(cell.value or "")) for cell in column) + 2))


@router.get("/runs/{schedule_run_id}/export")
def export_schedule(schedule_run_id: int, db: Session = Depends(get_db)) -> StreamingResponse:
    run = db.get(ScheduleRun, schedule_run_id)
    if not run:
        raise HTTPException(status_code=404, detail="생산 스케줄 이력을 찾을 수 없습니다.")
    if run.status != "확정":
        raise HTTPException(status_code=400, detail="Excel 다운로드는 확정된 생산계획 버전에서만 할 수 있습니다.")
    items = list(db.scalars(select(ProductionScheduleItem).where(ProductionScheduleItem.schedule_run_id == schedule_run_id).order_by(ProductionScheduleItem.planned_date, ProductionScheduleItem.line_code)))
    shortages = list(db.scalars(select(UnscheduledRequirement).where(UnscheduledRequirement.schedule_run_id == schedule_run_id).order_by(UnscheduledRequirement.product_code)))
    latest_validation = db.scalar(select(RawMaterialValidationRun).where(RawMaterialValidationRun.schedule_run_id == schedule_run_id).order_by(RawMaterialValidationRun.created_at.desc()))
    balances = list(db.scalars(select(RawMaterialDailyBalance).where(RawMaterialDailyBalance.validation_run_id == latest_validation.id).order_by(RawMaterialDailyBalance.balance_date, RawMaterialDailyBalance.plant_name, RawMaterialDailyBalance.material_code))) if latest_validation else []
    workbook = Workbook()
    summary = workbook.active
    _write_sheet(summary, "요약", ["제품", "필요 생산량(t)", "배정량(t)", "미배정량(t)"], [[row.product_code, row.required_quantity_ton, row.scheduled_quantity_ton, row.unallocated_quantity_ton] for row in shortages])
    schedule_sheet = workbook.create_sheet()
    _write_sheet(schedule_sheet, "생산 스케줄", ["일자", "공장", "라인", "제품", "생산계획(t)", "가용능력(t)", "정비휴지(h)", "전환·세척(h)", "고정", "변경 사유"], [[row.planned_date, row.plant_name, row.line_name, row.product_code, row.planned_quantity_ton, row.available_capacity_ton, row.downtime_hours, row.changeover_hours, "Y" if row.is_locked else "", row.adjustment_note] for row in items])
    if balances:
        material_sheet = workbook.create_sheet()
        _write_sheet(material_sheet, "원료 검증", ["일자", "공장", "원료", "기초재고(t)", "입고(t)", "BOM 소요량(t)", "기말재고(t)", "부족량(t)"], [[row.balance_date, row.plant_name, row.material_code, row.opening_quantity_ton, row.inbound_quantity_ton, row.required_quantity_ton, row.ending_quantity_ton, row.shortage_quantity_ton] for row in balances])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    headers = {"Content-Disposition": f'attachment; filename="production_schedule_{schedule_run_id}.xlsx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)
